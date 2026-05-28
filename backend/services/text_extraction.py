from __future__ import annotations

import json
from pathlib import Path
from typing import Optional
from xml.etree import ElementTree

import pandas as pd
from docx import Document
from PIL import Image
from pypdf import PdfReader
from pptx import Presentation
from striprtf.striprtf import rtf_to_text
import pytesseract


def extract_text(file_path: str) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()

    try:
        if suffix == ".pdf":
            return _extract_pdf(path)
        if suffix == ".docx":
            return _extract_docx(path)
        if suffix == ".rtf":
            return _extract_rtf(path)
        if suffix == ".pptx":
            return _extract_pptx(path)
        if suffix in {".xlsx", ".xls"}:
            return _extract_excel(path)
        if suffix in {".csv", ".tsv"}:
            return _extract_csv(path, suffix)
        if suffix == ".json":
            return _extract_json(path)
        if suffix == ".xml":
            return _extract_xml(path)
        if suffix in {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}:
            return _extract_image(path)

        return path.read_text(encoding="utf-8", errors="ignore")
    except Exception:
        return ""


def extract_pdf_pages(file_path: str) -> list[str]:
    path = Path(file_path)
    reader = PdfReader(str(path))
    pages: list[str] = []
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return pages


def _extract_pdf(path: Path) -> str:
    return "\n".join(extract_pdf_pages(str(path)))


def _extract_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text)


def _extract_rtf(path: Path) -> str:
    return rtf_to_text(path.read_text(encoding="utf-8", errors="ignore"))


def _extract_pptx(path: Path) -> str:
    presentation = Presentation(str(path))
    parts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
    return "\n".join(parts)


def _extract_excel(path: Path) -> str:
    data = pd.read_excel(path, sheet_name=None, engine="openpyxl")
    parts = []
    for name, frame in data.items():
        parts.append(f"# Sheet: {name}\n{frame.to_csv(index=False)}")
    return "\n".join(parts)


def _extract_csv(path: Path, suffix: str) -> str:
    sep = "\t" if suffix == ".tsv" else ","
    frame = pd.read_csv(path, sep=sep)
    return frame.to_csv(index=False)


def _extract_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return json.dumps(data, ensure_ascii=False, indent=2)


def _extract_xml(path: Path) -> str:
    tree = ElementTree.parse(path)
    return ElementTree.tostring(tree.getroot(), encoding="unicode")


def _extract_image(path: Path) -> str:
    image = Image.open(path)
    return pytesseract.image_to_string(image)
